import pandas as pd 
from docx import Document
doc = Document('5.報名表正面.docx')
df= pd.read_excel('1.中壢高商(14901).xlsx', sheet_name='Data-全測', skiprows=2)
rows = df.shape[0]
table = doc.tables[0]
nowcommend = ''
for i in range(0,rows):
    testset= set()
    for idxr, row in enumerate(table.rows):
        for  idxc , cell in enumerate(row.cells):
            if cell.text!='':
                if cell.text not in testset:
                    if nowcommend == '聯絡電話':
                        if cell.text not in testset:
                            homnnum,clphone = cell.text.split('\n')
                            homnnum = homnnum.replace(' ','')
                            clphone = clphone.replace(' ','')
                            homnnum = homnnum + str(df.loc[i,'電話(公)'])
                            clphone = clphone + '0' +str(df.loc[i,'電話(行動)'])
                            cell.text = homnnum + '\n' + clphone
                            testset.add(cell.text)
                    elif nowcommend =='年級':
                        reCell  = table.cell(idxr+1 , idxc-1)
                        reCell.text = str(df.loc[i,'年級'])
                    elif nowcommend =='班別':
                        reCell  = table.cell(idxr+1 , idxc-1)
                        reCell.text = str(df.loc[i, '班別'])
                    
                    testset.add(cell.text)
                    nowcommend = cell.text
                    nowcommend = nowcommend.replace('\n','')
    new_file_path = str(i+1)+'.docx'
    doc.save(new_file_path) 